"""Resume file validation and text extraction (PDF / DOCX)."""
import io
import re

MAX_SIZE_BYTES = 5 * 1024 * 1024
ALLOWED_EXTENSIONS = (".pdf", ".docx")


class ResumeParseError(ValueError):
    """Raised for invalid/unsupported/corrupt resume files. `status` is the HTTP code to use."""

    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text)
    return text.strip()


def _read_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise ResumeParseError("The PDF is password-protected. Please upload an unlocked copy.", 422)
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError(f"Could not read the PDF file (it may be corrupted): {exc}", 422)


def _read_docx(content: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        raise ResumeParseError(f"Could not read the DOCX file (it may be corrupted): {exc}", 422)


def extract_text(filename: str, content: bytes) -> str:
    name = (filename or "").lower()
    if not name.endswith(ALLOWED_EXTENSIONS):
        raise ResumeParseError("Unsupported file type. Please upload a PDF or DOCX resume.", 415)
    if not content:
        raise ResumeParseError("The uploaded file is empty.", 400)
    if len(content) > MAX_SIZE_BYTES:
        raise ResumeParseError("File is too large. Maximum allowed size is 5 MB.", 413)

    if name.endswith(".pdf"):
        if not content.startswith(b"%PDF"):
            raise ResumeParseError("This file has a .pdf extension but is not a valid PDF.", 422)
        text = _read_pdf(content)
    else:
        if not content.startswith(b"PK"):
            raise ResumeParseError("This file has a .docx extension but is not a valid DOCX.", 422)
        text = _read_docx(content)

    text = _clean(text)
    if len(text) < 50:
        raise ResumeParseError(
            "Almost no text could be extracted. If this is a scanned image PDF, please upload a text-based resume.",
            422,
        )
    return text
